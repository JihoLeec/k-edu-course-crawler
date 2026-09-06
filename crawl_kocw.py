#!/usr/bin/env python3
"""KOCW 대학강의·기관강의 크롤러 — KOCW 시트 양식으로 엑셀 출력."""

from __future__ import annotations

import argparse
import json
import logging
import re
import shutil
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterator
from urllib.parse import urljoin

import openpyxl
import requests
from bs4 import BeautifulSoup

BASE_URL = "https://www.kocw.net/home/search/"
DETAIL_URL = "https://www.kocw.net/home/cview.do?cid={cid}"
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
        "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "ko-KR,ko;q=0.9",
}
EXCEL_COLUMNS = (
    "순번",
    "출처(제작사)",
    "대메뉴",
    "중분류",
    "세분류(전공분류)",
    "상태",
    "콘텐츠명",
    "강좌수",
    "강의목차",
    "강사명",
    "URL",
    "비고",
)
TEMPLATE_FILE = "제발정신차려 이걸 날리면 어떡하니.xlsx"
SHEET_NAME = "KOCW"

log = logging.getLogger(__name__)


@dataclass
class CourseStub:
    cid: str
    title: str
    provider: str
    professor: str
    menu: str  # 대학강의 | 기관강의
    url: str


@dataclass
class CourseRow:
    provider: str
    menu: str
    mid_category: str
    sub_category: str
    title: str
    lecture_count: int
    syllabus: str
    professor: str
    url: str
    detail_fetched: bool = True


class KocwClient:
    def __init__(self, delay: float = 0.3, retries: int = 3):
        self.session = requests.Session()
        self.session.headers.update(HEADERS)
        self.delay = delay
        self.retries = retries

    def get(self, url: str) -> str:
        last_error: Exception | None = None
        for attempt in range(1, self.retries + 1):
            try:
                time.sleep(self.delay)
                resp = self.session.get(url, timeout=30)
                resp.raise_for_status()
                resp.encoding = "utf-8"
                return resp.text
            except Exception as exc:  # noqa: BLE001
                last_error = exc
                log.warning("요청 실패 (%s/%s): %s — %s", attempt, self.retries, url, exc)
                time.sleep(attempt)
        raise RuntimeError(f"요청 실패: {url}") from last_error


def parse_topic(topic: str) -> tuple[str, str]:
    parts = [p.strip() for p in re.split(r"\s*>\s*", topic) if p.strip()]
    if not parts:
        return "", ""
    if len(parts) == 1:
        return parts[0], ""
    return parts[0], " >".join(parts[1:])


def parse_syllabus(soup: BeautifulSoup) -> tuple[int, str]:
    items: list[str] = []
    for tr in soup.select("table tr"):
        tds = tr.select("td")
        if len(tds) < 3:
            continue
        num = tds[0].get_text(strip=True)
        name = tds[2].get_text(strip=True)
        if not re.match(r"^\d+\.?$", num) or not name:
            continue
        num_clean = num if num.endswith(".") else f"{num}."
        items.append(f"{num_clean}\t{name}")
    return len(items), "\n".join(items)


def parse_course_detail(html: str, stub: CourseStub) -> CourseRow:
    soup = BeautifulSoup(html, "lxml")

    title_el = soup.select_one(".detailTitle a")
    title = title_el.get_text(strip=True) if title_el else stub.title

    infos = [li.get_text(strip=True) for li in soup.select(".detailTitInfo li")]
    provider = infos[0] if infos else stub.provider
    professor = infos[1].strip() if len(infos) > 1 else stub.professor.strip()

    topic = ""
    for dl in soup.select(".detailViewList dl"):
        dt = dl.select_one("dt")
        if dt and dt.get_text(strip=True) == "주제분류":
            dd = dl.select_one("dd")
            topic = dd.get_text(" ", strip=True) if dd else ""
            break

    mid_category, sub_category = parse_topic(topic)
    lecture_count, syllabus = parse_syllabus(soup)

    return CourseRow(
        provider=provider or stub.provider,
        menu=stub.menu,
        mid_category=mid_category,
        sub_category=sub_category,
        title=title or stub.title,
        lecture_count=lecture_count,
        syllabus=syllabus,
        professor=professor or stub.professor,
        url=stub.url,
        detail_fetched=True,
    )


def stub_to_row(stub: CourseStub) -> CourseRow:
    """상세 페이지 수집 실패 시 목록에서 확보한 기본 정보만 반환."""
    return CourseRow(
        provider=stub.provider,
        menu=stub.menu,
        mid_category="",
        sub_category="",
        title=stub.title,
        lecture_count=0,
        syllabus="",
        professor=stub.professor.strip(),
        url=stub.url,
        detail_fetched=False,
    )


def cache_entry_fetched(entry: dict) -> bool:
    return entry.get("detail_fetched", True)


def row_from_cache(entry: dict) -> CourseRow:
    data = {k: v for k, v in entry.items() if k != "detail_fetched"}
    data.setdefault("detail_fetched", entry.get("detail_fetched", True))
    return CourseRow(**data)


def parse_listing_item(li: BeautifulSoup, menu: str) -> CourseStub | None:
    title_el = li.select_one(".listCon2 dt a")
    if not title_el:
        return None

    href = title_el.get("href", "")
    m = re.search(r"cid=([0-9a-f]+)", href)
    if not m:
        return None

    spans = li.select(".writer span")
    provider = spans[0].get_text(strip=True) if spans else ""
    professor = spans[1].get_text(strip=True) if len(spans) > 1 else provider

    return CourseStub(
        cid=m.group(1),
        title=title_el.get_text(strip=True),
        provider=provider,
        professor=professor,
        menu=menu,
        url=urljoin("https://www.kocw.net", href),
    )


def parse_max_page(soup: BeautifulSoup) -> int:
    pages = []
    for a in soup.select(".paging_area a"):
        href = a.get("href", "")
        m = re.search(r"page=(\d+)", href)
        if m:
            pages.append(int(m.group(1)))
        if a.get_text(strip=True).isdigit():
            pages.append(int(a.get_text(strip=True)))
    return max(pages) if pages else 1


def parse_univ_providers(html: str) -> list[tuple[str, str, str]]:
    soup = BeautifulSoup(html, "lxml")
    providers: list[tuple[str, str, str]] = []
    for a in soup.select('a[id^="univ_"]'):
        if a.get("id") == "univ_0":
            continue
        href = a.get("href", "")
        ud = re.search(r"ud=(\d+)", href)
        class_id = a.get("id", "")
        name = a.get_text(" ", strip=True)
        name = re.sub(r"\s*\(\d+\)\s*$", "", name)
        if ud and class_id:
            providers.append((name, ud.group(1), class_id))
    return providers


def parse_org_providers(html: str) -> list[tuple[str, str, str]]:
    soup = BeautifulSoup(html, "lxml")
    providers: list[tuple[str, str, str]] = []
    for a in soup.select('a[href*="orgCoursesAll.do?od="]'):
        href = a.get("href", "")
        od = re.search(r"od=(\d+)", href)
        class_id = re.search(r"classId=(org_\d+)", href)
        name = a.get_text(" ", strip=True)
        name = re.sub(r"\s*\(\d+\)\s*$", "", name)
        if od and class_id:
            providers.append((name, od.group(1), class_id.group(1)))
    return providers


def iter_listing_pages(
    client: KocwClient,
    menu: str,
    providers: list[tuple[str, str, str]],
    build_url,
    limit: int = 0,
) -> Iterator[CourseStub]:
    seen: set[str] = set()
    count = 0
    for name, provider_id, class_id in providers:
        page = 1
        max_page = 1
        while page <= max_page:
            html = client.get(build_url(provider_id, class_id, page))
            soup = BeautifulSoup(html, "lxml")
            max_page = parse_max_page(soup)
            if page == 1 or page % 10 == 0 or page == max_page:
                log.info("%s 목록: %s %s/%s페이지", menu, name, page, max_page)
            for li in soup.select(".lectContList > li"):
                stub = parse_listing_item(li, menu)
                if stub and stub.cid not in seen:
                    seen.add(stub.cid)
                    yield stub
                    count += 1
                    if limit > 0 and count >= limit:
                        return
            page += 1


def load_cache(path: Path) -> dict[str, dict]:
    if not path.exists():
        return {}
    with path.open(encoding="utf-8") as f:
        return json.load(f)


def save_stubs(path: Path, stubs: list[CourseStub], menus: set[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "menus": sorted(menus),
        "collected_at": datetime.now().isoformat(timespec="seconds"),
        "stubs": [asdict(s) for s in stubs],
    }
    with path.open("w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)


def load_stubs(path: Path, menus: set[str]) -> list[CourseStub] | None:
    if not path.exists():
        return None
    with path.open(encoding="utf-8") as f:
        payload = json.load(f)
    stubs = [CourseStub(**item) for item in payload.get("stubs", [])]
    if menus:
        stubs = [s for s in stubs if s.menu in menus]
    if not stubs:
        return None
    log.info(
        "목록 캐시 사용: %s건 (%s, 수집일 %s)",
        len(stubs),
        path,
        payload.get("collected_at", "알 수 없음"),
    )
    return stubs


def get_stubs(
    client: KocwClient,
    menus: set[str],
    stubs_path: Path,
    limit: int,
    refresh_listing: bool,
) -> list[CourseStub]:
    if limit == 0 and not refresh_listing:
        cached = load_stubs(stubs_path, menus)
        if cached:
            return cached

    stubs = collect_stubs(client, menus, limit=limit)
    if limit == 0:
        save_stubs(stubs_path, stubs, menus)
        log.info("목록 캐시 저장: %s", stubs_path)
    return stubs
    if not path.exists():
        return {}
    with path.open(encoding="utf-8") as f:
        return json.load(f)


def save_cache(path: Path, cache: dict[str, dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False, indent=2)


def fetch_details(
    client: KocwClient,
    stubs: list[CourseStub],
    cache_path: Path,
    workers: int,
) -> list[CourseRow]:
    cache = load_cache(cache_path)
    pending: list[CourseStub] = []

    for stub in stubs:
        entry = cache.get(stub.cid)
        if entry is None or not cache_entry_fetched(entry):
            pending.append(stub)

    cached_ok = len(stubs) - len(pending)
    log.info("상세 페이지: 캐시 %s건, 신규 %s건", cached_ok, len(pending))

    def fetch_one(stub: CourseStub) -> tuple[str, CourseRow, bool]:
        try:
            html = client.get(DETAIL_URL.format(cid=stub.cid))
            return stub.cid, parse_course_detail(html, stub), True
        except RuntimeError:
            log.error("상세 수집 실패(목록 정보만 사용): %s", stub.url)
            return stub.cid, stub_to_row(stub), False

    failed = 0
    if pending:
        with ThreadPoolExecutor(max_workers=workers) as executor:
            futures = {executor.submit(fetch_one, stub): stub for stub in pending}
            done = 0
            for future in as_completed(futures):
                cid, row, ok = future.result()
                if ok:
                    cache[cid] = {**asdict(row), "detail_fetched": True}
                else:
                    failed += 1
                done += 1
                if done % 50 == 0:
                    save_cache(cache_path, cache)
                    log.info("상세 수집 진행: %s/%s", done, len(pending))
        save_cache(cache_path, cache)

    if failed:
        log.warning("상세 수집 실패 %s건 — 엑셀에는 기본 정보만 기록, 재실행 시 재시도", failed)

    rows: list[CourseRow] = []
    for stub in stubs:
        entry = cache.get(stub.cid)
        if entry and cache_entry_fetched(entry):
            rows.append(row_from_cache(entry))
        else:
            rows.append(stub_to_row(stub))

    return rows


def _append_rows(ws, rows: list[CourseRow]) -> None:
    for idx, row in enumerate(rows, start=1):
        ws.append(
            [
                idx,
                row.provider,
                row.menu,
                row.mid_category,
                row.sub_category,
                None,
                row.title,
                row.lecture_count,
                row.syllabus,
                row.professor,
                row.url,
                None if row.detail_fetched else "상세수집실패",
            ]
        )


def write_excel(template: Path | None, output: Path, rows: list[CourseRow]) -> None:
    if template and template.exists():
        shutil.copy2(template, output)
        wb = openpyxl.load_workbook(output)
        if SHEET_NAME in wb.sheetnames:
            ws = wb[SHEET_NAME]
            if ws.max_row > 1:
                ws.delete_rows(2, ws.max_row - 1)
        else:
            ws = wb.create_sheet(SHEET_NAME, 0)
            ws.append(list(EXCEL_COLUMNS))
        log.info("템플릿 사용: %s", template)
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = SHEET_NAME
        ws.append(list(EXCEL_COLUMNS))
        log.info("템플릿 없음 — KOCW 시트만 새로 생성")

    _append_rows(ws, rows)
    wb.save(output)
    wb.close()


def collect_stubs(client: KocwClient, menus: set[str], limit: int = 0) -> list[CourseStub]:
    stubs: list[CourseStub] = []
    remaining = limit

    if "대학강의" in menus:
        log.info("대학 목록 수집 중...")
        html = client.get(urljoin(BASE_URL, "univCoursesAll.do"))
        providers = parse_univ_providers(html)
        log.info("대학 %s곳", len(providers))
        for stub in iter_listing_pages(
            client,
            "대학강의",
            providers,
            lambda ud, class_id, page: (
                f"{BASE_URL}univCoursesAll.do?ud={ud}&so1=1&so2=2"
                f"&classId1=char_1&classId2={class_id}&page={page}"
            ),
            limit=remaining,
        ):
            stubs.append(stub)
            if limit > 0:
                remaining = limit - len(stubs)
                if remaining <= 0:
                    return stubs

    if "기관강의" in menus:
        log.info("기관 목록 수집 중...")
        html = client.get(urljoin(BASE_URL, "orgCoursesAll.do"))
        providers = parse_org_providers(html)
        log.info("기관 %s곳", len(providers))
        for stub in iter_listing_pages(
            client,
            "기관강의",
            providers,
            lambda od, class_id, page: (
                f"{BASE_URL}orgCoursesAll.do?od={od}&so1=1&so2=2"
                f"&classId={class_id}&page={page}"
            ),
            limit=remaining,
        ):
            stubs.append(stub)
            if limit > 0:
                remaining = limit - len(stubs)
                if remaining <= 0:
                    return stubs

    return stubs


def main() -> None:
    parser = argparse.ArgumentParser(description="KOCW 강의 크롤러")
    parser.add_argument(
        "--template",
        default=TEMPLATE_FILE,
        help="양식 엑셀 파일 경로 (없으면 KOCW 시트만 새로 생성)",
    )
    parser.add_argument(
        "--output",
        default=None,
        help="결과 엑셀 파일 경로 (기본: KOCW_크롤링결과_날짜.xlsx)",
    )
    parser.add_argument(
        "--menu",
        choices=["all", "univ", "org"],
        default="all",
        help="크롤링 대상 (all=대학+기관)",
    )
    parser.add_argument("--limit", type=int, default=0, help="테스트용 수집 건수 제한")
    parser.add_argument("--delay", type=float, default=0.3, help="요청 간격(초)")
    parser.add_argument("--workers", type=int, default=5, help="상세 페이지 병렬 수")
    parser.add_argument(
        "--cache",
        default=".cache/kocw_details.json",
        help="상세 페이지 캐시 파일",
    )
    parser.add_argument(
        "--stubs-cache",
        default=".cache/kocw_stubs.json",
        help="강의 목록 캐시 파일",
    )
    parser.add_argument(
        "--refresh-listing",
        action="store_true",
        help="저장된 목록 캐시를 무시하고 목록을 다시 수집",
    )
    parser.add_argument("-v", "--verbose", action="store_true")
    args = parser.parse_args()

    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
    )

    template_path = Path(args.template)
    template: Path | None = template_path if template_path.exists() else None
    if template is None:
        log.warning("템플릿 없음 (%s) — KOCW 시트만 생성합니다", template_path)

    output = Path(
        args.output
        or f"KOCW_크롤링결과_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    )
    cache_path = Path(args.cache)
    stubs_path = Path(args.stubs_cache)

    menus = {"대학강의", "기관강의"}
    if args.menu == "univ":
        menus = {"대학강의"}
    elif args.menu == "org":
        menus = {"기관강의"}

    client = KocwClient(delay=args.delay)
    stubs = get_stubs(
        client,
        menus,
        stubs_path,
        limit=args.limit,
        refresh_listing=args.refresh_listing,
    )
    log.info("목록 준비 완료: %s건", len(stubs))

    rows = fetch_details(client, stubs, cache_path, args.workers)
    write_excel(template, output, rows)
    log.info("저장 완료: %s (%s건)", output, len(rows))


if __name__ == "__main__":
    main()
